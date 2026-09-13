import csv
import io
import logging
import os
import secrets
import sys
from datetime import datetime, timedelta

import jwt
from fastapi import FastAPI, Query, HTTPException, Depends, Header, Request
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response, JSONResponse
from fastapi.staticfiles import StaticFiles
from passlib.hash import argon2
from pydantic import BaseModel, Field

from database import (
    get_stations,
    get_station_by_id,
    insert_water_data,
    get_water_history,
    get_history_since,
    get_station_stats,
    count_warnings_since,
    get_all_latest_water_data,
    get_rainfall_summary,
    get_warnings,
    get_active_warnings,
    get_warnings_summary,
    handle_warning,
    get_latest_warning,
    get_all_history_since,
    count_all_warnings_since,
    get_report_data,
    get_report_stats,
    get_user_by_username,
    get_user_by_id,
    get_all_users,
    create_user,
    delete_user,
    update_user_role,
    get_alerts,
    count_alerts,
    get_alert_summary,
    acknowledge_alert,
    resolve_alert,
    get_data_quality_stats,
)
from services.data_analysis import (
    analyze_trend,
    forecast,
    build_comparison,
    DEFAULT_FORECAST_HORIZON,
)
from services.alert_service import create_alert_if_needed
from services.ai_analysis_service import generate_ai_analysis, get_ai_analyzer
from services.rate_limiter import allow_ai_request
from services.water_data_provider import get_provider, ProviderError
from services.real_water_provider import RealWaterProvider
from services.mock_water_provider import build_mock_record
from services.data_normalizer import (
    SOURCE_CHENGDU_OPEN_DATA,
    SOURCE_MOCK_FALLBACK,
    QUALITY_DEGRADED,
    QUALITY_VALID,
)

logger = logging.getLogger("water_monitor")

JWT_SECRET_KEY = os.environ.get("JWT_SECRET_KEY", "")
if not JWT_SECRET_KEY:
    JWT_SECRET_KEY = secrets.token_hex(32)
    print(
        "WARNING: 未设置 JWT_SECRET_KEY，已生成随机密钥；重启后已签发令牌将全部失效。"
        "生产环境请通过环境变量设置稳定的 JWT_SECRET_KEY。",
        file=sys.stderr,
    )
JWT_ALGORITHM = "HS256"
JWT_EXPIRE_MINUTES = int(os.environ.get("JWT_EXPIRE_MINUTES", "60"))

app = FastAPI()

_cors_origins = os.environ.get("CORS_ORIGINS", "").strip()
if _cors_origins:
    allow_origins = [o.strip() for o in _cors_origins.split(",") if o.strip()]
else:
    allow_origins = [
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ]

app.add_middleware(
    CORSMiddleware,
    allow_origins=allow_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["*"],
)

# ── 第21阶段：安全响应头（应用级中间件，统一注入到所有响应）──
_CSP_POLICY = (
    "default-src 'self'; "
    "script-src 'self'; "
    "style-src 'self' 'unsafe-inline'; "
    "img-src 'self' data: https://*.tile.openstreetmap.org; "
    "font-src 'self' data:; "
    "connect-src 'self' http://127.0.0.1:8000 http://localhost:8000 "
    "ws://localhost:5173 ws://127.0.0.1:5173; "
    "object-src 'none'; base-uri 'self'; frame-ancestors 'self';"
)


@app.middleware("http")
async def add_security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    if os.environ.get("ENABLE_CSP", "1") not in ("0", "false", "False"):
        response.headers.setdefault("Content-Security-Policy", _CSP_POLICY)
    return response


# ── 第21阶段：统一异常处理，避免向客户端泄露内部实现细节 ──
@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    return JSONResponse(
        status_code=exc.status_code,
        content={"detail": exc.detail},
        headers=exc.headers,
    )


@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    return JSONResponse(status_code=422, content={"detail": exc.errors()})


@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception):
    logger.exception("未处理异常: method=%s path=%s", request.method, request.url.path)
    return JSONResponse(status_code=500, content={"detail": "服务器内部错误。请稍后重试。"})


_dist_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "frontend", "dist")


class LoginRequest(BaseModel):
    username: str = Field(max_length=64)
    password: str = Field(max_length=200)


class RegisterRequest(BaseModel):
    username: str = Field(max_length=30)
    password: str = Field(max_length=200)
    confirm_password: str = Field(default="", max_length=200)


class CreateUserRequest(BaseModel):
    username: str = Field(max_length=30)
    password: str = Field(max_length=200)
    role: str = "user"


class UpdateRoleRequest(BaseModel):
    role: str = Field(min_length=1, max_length=32)


def create_access_token(data: dict) -> str:
    to_encode = data.copy()
    expire = datetime.utcnow() + timedelta(minutes=JWT_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, JWT_SECRET_KEY, algorithm=JWT_ALGORITHM)


def get_current_user(authorization: str = Header(default="")):
    if not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="未登录")
    token = authorization[7:]
    try:
        payload = jwt.decode(token, JWT_SECRET_KEY, algorithms=[JWT_ALGORITHM])
        user_id = payload.get("user_id")
        if user_id is None:
            raise HTTPException(status_code=401, detail="无效的token")
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="token已过期")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="无效的token")
    user = get_user_by_id(user_id)
    if not user or not user["is_active"]:
        raise HTTPException(status_code=401, detail="用户不存在或已禁用")
    return user


def require_admin(user: dict = Depends(get_current_user)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="需要管理员权限")
    return user


# 防时序探测：用户不存在时也执行一次 Argon2 校验，避免通过响应时间猜测用户名。
_DUMMY_PASSWORD_HASH = argon2.hash("timing-equalizer-dummy-password")


# ──────────────────────────── Auth API ────────────────────────────

@app.post("/api/auth/login")
def login(req: LoginRequest):
    user = get_user_by_username(req.username)
    if not user:
        try:
            argon2.verify(req.password, _DUMMY_PASSWORD_HASH)
        except Exception:
            pass
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    if not user["is_active"]:
        # 禁用账号与“用户不存在/密码错误”返回一致，不泄露账号是否存在
        try:
            argon2.verify(req.password, user["password_hash"])
        except Exception:
            pass
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    if not argon2.verify(req.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    token = create_access_token({"user_id": user["id"], "username": user["username"], "role": user["role"]})
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": user["id"],
            "username": user["username"],
            "role": user["role"],
        },
    }


@app.post("/api/auth/register")
def register(req: RegisterRequest):
    username = (req.username or "").strip()
    password = req.password or ""

    if not username:
        raise HTTPException(status_code=400, detail="用户名不能为空")
    if len(username) < 3 or len(username) > 30:
        raise HTTPException(status_code=400, detail="用户名长度需在3到30个字符之间")
    if len(password) < 8:
        raise HTTPException(status_code=400, detail="密码长度不能少于8位")
    if req.confirm_password and req.confirm_password != password:
        raise HTTPException(status_code=400, detail="两次密码输入不一致")

    if get_user_by_username(username):
        raise HTTPException(status_code=400, detail="用户名已存在")

    password_hash = argon2.hash(password)
    # 自主注册用户强制为普通用户 user，客户端传入的任何 role 都会被忽略
    user_id = create_user(username=username, password_hash=password_hash, role="user")
    return {"success": True, "message": "注册成功", "user_id": user_id}


@app.get("/api/auth/me")
def get_me(user: dict = Depends(get_current_user)):
    return {
        "id": user["id"],
        "username": user["username"],
        "role": user["role"],
        "is_active": user["is_active"],
        "created_at": user["created_at"],
    }


@app.post("/api/auth/logout")
def logout():
    return {"success": True}


# ──────────────────────────── User Management API ────────────────────────────

@app.get("/api/users")
def list_users(admin: dict = Depends(require_admin)):
    users = get_all_users()
    return {"data": users}


@app.post("/api/users")
def create_user_endpoint(req: CreateUserRequest, admin: dict = Depends(require_admin)):
    existing = get_user_by_username(req.username)
    if existing:
        raise HTTPException(status_code=400, detail="用户名已存在")
    if req.role not in ("admin", "user"):
        raise HTTPException(status_code=400, detail="角色必须是 admin 或 user")
    password_hash = argon2.hash(req.password)
    user_id = create_user(username=req.username, password_hash=password_hash, role=req.role)
    return {"success": True, "user_id": user_id}


@app.delete("/api/users/{user_id}")
def delete_user_endpoint(user_id: int, admin: dict = Depends(require_admin)):
    if user_id == admin["id"]:
        raise HTTPException(status_code=400, detail="不能删除自己")
    success = delete_user(user_id)
    if not success:
        raise HTTPException(status_code=404, detail="用户不存在")
    return {"success": True}


@app.post("/api/users/{user_id}/role")
def update_user_role_endpoint(user_id: int, req: UpdateRoleRequest, admin: dict = Depends(require_admin)):
    if req.role not in ("admin", "user"):
        raise HTTPException(status_code=400, detail="角色必须是 admin 或 user")
    if user_id == admin["id"]:
        raise HTTPException(status_code=400, detail="不能修改自己的角色")
    success = update_user_role(user_id, req.role)
    if not success:
        raise HTTPException(status_code=404, detail="用户不存在")
    return {"success": True}


# ──────────────────────────── Original Monitoring API (unchanged) ────────────────────────────

@app.get("/api/health")
def health_check():
    return {
        "status": "ok",
        "message": "智水监测后端运行正常",
    }


@app.get("/api/stations")
def stations():
    data = get_stations()
    return {"data": data}


def _build_current_record(station: dict) -> dict:
    """构造一条站点当前记录。

    优先使用成都官方雨情 API 提供的真实降雨量；
    官方数据不可用或未配置时回退到模拟数据，并按约定标注数据来源与质量。
    """
    provider = get_provider()
    if provider.name == SOURCE_CHENGDU_OPEN_DATA:
        try:
            record = provider.get_station_data(station["station_id"])
            if record:
                return record
        except ProviderError:
            pass
        except Exception:
            pass
        record = build_mock_record(station)
        record["source"] = SOURCE_MOCK_FALLBACK
        record["data_quality"] = QUALITY_DEGRADED
        return record

    record = build_mock_record(station)
    record["source"] = "mock"
    record["data_quality"] = QUALITY_VALID
    return record


@app.get("/api/water-data")
def get_water_data(station_id: str = Query(default="ST001", max_length=40)):
    station = get_station_by_id(station_id)
    if not station:
        return {"error": f"水文站 {station_id} 不存在"}

    record = _build_current_record(station)

    insert_water_data(
        station_id=station_id,
        station_name=station["station_name"],
        water_level=record["water_level"],
        warning_level=record["warning_level"],
        rainfall=record["rainfall"],
        status=record["status"],
        source=record["source"],
        data_quality=record["data_quality"],
        created_at=record["timestamp"],
    )

    if record["status"] != "正常":
        create_alert_if_needed(
            station_id=station_id,
            station_name=station["station_name"],
            water_level=record["water_level"],
            warning_level=record["warning_level"],
            rainfall=record["rainfall"],
            status=record["status"],
        )

    return {
        "station_id": station_id,
        "station_name": station["station_name"],
        "water_level": record["water_level"],
        "warning_level": record["warning_level"],
        "rainfall": record["rainfall"],
        "status": record["status"],
        "source": record["source"],
        "data_quality": record["data_quality"],
        "updated_at": record["timestamp"],
    }


@app.get("/api/water-data-all")
def get_water_data_all():
    data = get_all_latest_water_data()
    return {"data": data}


@app.get("/api/data-source")
def data_source():
    provider = get_provider()
    result = {
        "source": provider.name,
        "name": provider.name,
        "display_name": provider.display_name,
        "configured": provider.is_configured(),
    }
    if isinstance(provider, RealWaterProvider):
        status = provider.check_connection()
    else:
        status = {
            "provider": provider.name,
            "display_name": provider.display_name,
            "configured": False,
            "signature_ok": False,
            "reachable": False,
            "authenticated": False,
            "data_valid": False,
            "data_quality": QUALITY_DEGRADED,
            "tested_stcd": False,
            "reason": "未启用成都官方数据源（缺少 CHENGDU_CLIENT_ID / CHENGDU_CLIENT_SECRET 配置）",
        }
    result["status"] = status
    result["reachable"] = status["reachable"]
    result["data_quality"] = status["data_quality"]
    return result


@app.get("/api/data-quality")
def data_quality():
    stats = get_data_quality_stats()
    return {
        "source": get_provider().name,
        "data_quality": stats,
        "quality": stats,
    }


@app.get("/api/water-history")
def water_history(
    station_id: str = Query(default="ST001", max_length=40),
    limit: int = Query(default=20, ge=1, le=500),
):
    data = get_water_history(station_id=station_id, limit=limit)
    return {"data": data}


@app.get("/api/rainfall-summary")
def rainfall_summary(station_id: str = Query(default="ST001", max_length=40)):
    data = get_rainfall_summary(station_id=station_id)
    return data


@app.get("/api/warnings")
def warnings_list(limit: int = Query(default=50, ge=1, le=500)):
    data = get_warnings(limit=limit)
    return {"data": data}


@app.get("/api/warnings/active")
def warnings_active():
    data = get_active_warnings()
    return {"data": data}


@app.get("/api/warnings-summary")
def warnings_summary():
    return get_warnings_summary()


@app.post("/api/warnings/{warning_id}/handle")
def handle_warning_endpoint(warning_id: int, admin: dict = Depends(require_admin)):
    success = handle_warning(warning_id)
    if not success:
        raise HTTPException(status_code=404, detail=f"预警记录 {warning_id} 不存在")
    return {"success": True}


# ──────────────────────────── 第14阶段：历史数据分析与趋势预测 ────────────────────────────

HOURS_MIN = 1
HOURS_MAX = 720


def _since_cutoff(hours: int):
    return (datetime.now() - timedelta(hours=hours)).isoformat(timespec="seconds")


def _get_station_or_404(station_id: str):
    station = get_station_by_id(station_id)
    if not station:
        raise HTTPException(status_code=404, detail=f"水文站 {station_id} 不存在")
    return station


@app.get("/api/history")
def history_api(
    station_id: str = Query(default="ST001", max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    station = _get_station_or_404(station_id)
    since = _since_cutoff(hours)
    rows = get_history_since(station_id=station_id, since_iso=since)

    records = [
        {
            "timestamp": r["created_at"],
            "water_level": r["water_level"],
            "warning_level": r["warning_level"],
            "rainfall": r["rainfall"],
            "status": r["status"],
            "source": r["source"] or "mock",
            "data_quality": r["data_quality"] or "valid",
        }
        for r in rows
    ]

    return {
        "station_id": station_id,
        "station_name": station["station_name"],
        "hours": hours,
        "count": len(records),
        "data": records,
    }


@app.get("/api/trend")
def trend_api(
    station_id: str = Query(default="ST001", max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    station = _get_station_or_404(station_id)
    since = _since_cutoff(hours)
    rows = get_history_since(station_id=station_id, since_iso=since)

    result = analyze_trend(rows, hours=hours)
    return {
        "station_id": station_id,
        "station_name": station["station_name"],
        "hours": hours,
        **result,
    }


@app.get("/api/statistics")
def statistics_api(
    station_id: str = Query(default="ST001", max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    station = _get_station_or_404(station_id)
    since = _since_cutoff(hours)
    st = get_station_stats(station_id=station_id, since_iso=since)
    warning_count = count_warnings_since(station_id=station_id, since_iso=since)

    n = st["n"]
    data_points = int(n)
    if n <= 0:
        stats = {
            "max_water_level": None,
            "min_water_level": None,
            "avg_water_level": None,
            "max_rainfall": None,
            "avg_rainfall": None,
            "total_rainfall": 0.0,
            "warning_count": warning_count,
            "data_points": data_points,
        }
    else:
        stats = {
            "max_water_level": round(st["mx"], 2),
            "min_water_level": round(st["mn"], 2),
            "avg_water_level": round(st["av"], 2),
            "max_rainfall": round(st["mxr"], 1),
            "avg_rainfall": round(st["avr"], 2),
            "total_rainfall": round(st["tot"], 1),
            "warning_count": warning_count,
            "data_points": data_points,
        }

    return {
        "station_id": station_id,
        "station_name": station["station_name"],
        "hours": hours,
        "sufficient": data_points >= 1,
        "message": "" if data_points >= 1 else "历史数据不足",
        **stats,
    }


@app.get("/api/forecast")
def forecast_api(
    station_id: str = Query(default="ST001", max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
    horizon_hours: int = Query(default=DEFAULT_FORECAST_HORIZON, ge=1, le=72),
):
    station = _get_station_or_404(station_id)
    since = _since_cutoff(hours)
    rows = get_history_since(station_id=station_id, since_iso=since)

    result = forecast(rows, warning_level=station["warning_level"], hours=hours,
                      horizon_hours=horizon_hours)
    return {
        "station_id": station_id,
        "station_name": station["station_name"],
        "hours": hours,
        **result,
    }


# ──────────────────────────── 第15阶段：多站综合对比分析 ────────────────────────────

@app.get("/api/comparison")
def comparison_api(
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    stations = get_stations()
    since = _since_cutoff(hours)
    all_history = get_all_history_since(since)
    warning_counts = count_all_warnings_since(since)

    result = build_comparison(stations, all_history, warning_counts, hours=hours)
    return {
        "hours": hours,
        "count": result["station_count"],
        **result,
    }


# ──────────────────────────── 第16阶段：数据报表导出 ────────────────────────────

@app.get("/api/report")
def report_data_api(
    station_id: str = Query(default=None, max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
    limit: int = Query(default=None, ge=1, le=10000),
):
    if station_id and not get_station_by_id(station_id):
        raise HTTPException(status_code=404, detail=f"水文站 {station_id} 不存在")
    since = _since_cutoff(hours)
    records = get_report_data(station_id=station_id, since_iso=since, limit=limit)
    station_name = None
    if station_id:
        s = get_station_by_id(station_id)
        station_name = s["station_name"] if s else station_id
    return {
        "station_id": station_id,
        "station_name": station_name,
        "hours": hours,
        "count": len(records),
        "data": records,
    }


@app.get("/api/report/statistics")
def report_statistics_api(
    station_id: str = Query(default=None, max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    if station_id and not get_station_by_id(station_id):
        raise HTTPException(status_code=404, detail=f"水文站 {station_id} 不存在")
    since = _since_cutoff(hours)
    st = get_report_stats(station_id=station_id, since_iso=since)
    n = st["n"]
    data_points = int(n)
    if data_points <= 0:
        return {
            "station_id": station_id,
            "station_name": (get_station_by_id(station_id) or {}).get("station_name", station_id) if station_id else None,
            "hours": hours,
            "data_points": 0,
            "avg_water_level": None,
            "max_water_level": None,
            "min_water_level": None,
            "avg_rainfall": None,
            "max_rainfall": None,
            "warning_count": st["warning_count"],
        }
    return {
        "station_id": station_id,
        "station_name": (get_station_by_id(station_id) or {}).get("station_name", station_id) if station_id else None,
        "hours": hours,
        "data_points": data_points,
        "avg_water_level": round(st["av"], 2),
        "max_water_level": round(st["mx"], 2),
        "min_water_level": round(st["mn"], 2),
        "avg_rainfall": round(st["avr"], 2),
        "max_rainfall": round(st["mxr"], 1),
        "warning_count": st["warning_count"],
    }


@app.get("/api/report/export/csv")
def report_export_csv(
    station_id: str = Query(default=None, max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    if station_id and not get_station_by_id(station_id):
        raise HTTPException(status_code=404, detail=f"水文站 {station_id} 不存在")
    since = _since_cutoff(hours)
    records = get_report_data(station_id=station_id, since_iso=since)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["站点编号", "站点名称", "采集时间", "当前水位", "警戒水位", "降雨量", "状态", "数据来源", "数据质量"])
    for r in records:
        writer.writerow([
            r["station_id"], r["station_name"], r["timestamp"],
            r["water_level"], r["warning_level"], r["rainfall"],
            r["status"], r["source"], r["data_quality"],
        ])

    csv_bytes = ("\ufeff" + output.getvalue()).encode("utf-8")
    now_str = datetime.now().strftime("%Y-%m-%d")
    filename = f"zhishui_water_report_{now_str}.csv"
    return Response(
        content=csv_bytes,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/report/export/excel")
def report_export_excel(
    station_id: str = Query(default=None, max_length=40),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    if station_id and not get_station_by_id(station_id):
        raise HTTPException(status_code=404, detail=f"水文站 {station_id} 不存在")
    since = _since_cutoff(hours)
    records = get_report_data(station_id=station_id, since_iso=since)

    wb = Workbook()
    ws = wb.active
    ws.title = "水情报表"

    # Title row
    ws.merge_cells("A1:I1")
    ws["A1"] = "智慧水利 · 水情数据报表"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Info row
    station_name = (get_station_by_id(station_id) or {}).get("station_name", station_id) if station_id else "全部站点"
    ws.merge_cells("A2:I2")
    ws["A2"] = f"时间范围：最近 {hours} 小时  |  站点范围：{station_name}  |  生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(size=10, color="666666")
    ws["A2"].alignment = Alignment(horizontal="left")

    # Header row
    headers = ["站点编号", "站点名称", "采集时间", "当前水位", "警戒水位", "降雨量", "状态", "数据来源", "数据质量"]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E2EFF4", end_color="E2EFF4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Data rows
    for i, r in enumerate(records, 4):
        ws.cell(row=i, column=1, value=r["station_id"])
        ws.cell(row=i, column=2, value=r["station_name"])
        ws.cell(row=i, column=3, value=r["timestamp"])
        ws.cell(row=i, column=4, value=r["water_level"])
        ws.cell(row=i, column=5, value=r["warning_level"])
        ws.cell(row=i, column=6, value=r["rainfall"])
        ws.cell(row=i, column=7, value=r["status"])
        ws.cell(row=i, column=8, value=r["source"])
        ws.cell(row=i, column=9, value=r["data_quality"])

    # Auto-adjust column widths
    for col_idx, h in enumerate(headers, 1):
        max_length = len(h) * 2
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=4, max_row=3 + len(records)):
            for cell in row:
                if cell.value is not None:
                    cell_str = str(cell.value)
                    cjk = sum(1 for c in cell_str if "\u4e00" <= c <= "\u9fff")
                    max_length = max(max_length, len(cell_str) + cjk)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 4, 40)

    # Freeze header
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    now_str = datetime.now().strftime("%Y-%m-%d")
    filename = f"zhishui_water_report_{now_str}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ──────────────────────────── 第17阶段：智能预警中心 ────────────────────────────

ALERT_LEVELS = ("normal", "attention", "warning", "danger")
ALERT_STATUSES = ("pending", "acknowledged", "resolved")


@app.get("/api/alerts")
def alerts_api(
    station_id: str = Query(default=None, max_length=40),
    level: str = Query(default=None),
    status: str = Query(default=None),
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
):
    if station_id and not get_station_by_id(station_id):
        raise HTTPException(status_code=422, detail=f"水文站 {station_id} 不存在")
    if level and level not in ALERT_LEVELS:
        raise HTTPException(status_code=422, detail="level 必须是 normal/attention/warning/danger 之一")
    if status and status not in ALERT_STATUSES:
        raise HTTPException(status_code=422, detail="status 必须是 pending/acknowledged/resolved 之一")
    since = _since_cutoff(hours)
    items = get_alerts(
        station_id=station_id,
        alert_level=level,
        status=status,
        since_iso=since,
        limit=limit,
        offset=offset,
    )
    total = count_alerts(station_id=station_id, alert_level=level, status=status, since_iso=since)
    return {"items": items, "total": total}


@app.get("/api/alerts/summary")
def alerts_summary_api():
    return get_alert_summary()


@app.post("/api/alerts/{alert_id}/acknowledge")
def alerts_acknowledge(alert_id: int, admin: dict = Depends(require_admin)):
    success = acknowledge_alert(alert_id, admin["username"])
    if not success:
        raise HTTPException(status_code=404, detail=f"预警记录 {alert_id} 不存在")
    return {"success": True}


@app.post("/api/alerts/{alert_id}/resolve")
def alerts_resolve(alert_id: int, admin: dict = Depends(require_admin)):
    success = resolve_alert(alert_id, admin["username"])
    if not success:
        raise HTTPException(status_code=404, detail=f"预警记录 {alert_id} 不存在")
    return {"success": True}


# ──────────────────────────── 第20阶段：AI 智能水情分析 ────────────────────────────


def _ai_request_identity(request: Request) -> str:
    """限流身份：优先 JWT 中的 user_id，其次客户端 IP。"""
    authorization = request.headers.get("authorization", "")
    if authorization.startswith("Bearer "):
        try:
            payload = jwt.decode(authorization[7:], JWT_SECRET_KEY, algorithms=[JWT_ALGORITHM])
            user_id = payload.get("user_id")
            if user_id is not None:
                return f"user:{user_id}"
        except jwt.InvalidTokenError:
            pass
    host = request.client.host if request.client else "unknown"
    return f"ip:{host}"


@app.get("/api/ai-analysis")
def ai_analysis_api(
    request: Request,
    hours: int = Query(default=24, ge=HOURS_MIN, le=HOURS_MAX),
):
    """按需生成结构化 AI 水情分析（只读，不写库）。

    默认规则分析（rule_based）；配置 AI_ANALYZER=openai 后调用真实 AI 模型，
    任何失败都会自动降级为规则分析，绝不向客户端抛出 500。
    真实 AI 模式受内存滑动窗口限流（默认 10 次/分钟，AI_RATE_LIMIT_PER_MINUTE）。
    """
    try:
        if get_ai_analyzer().analysis_source == "ai_model":
            identity = _ai_request_identity(request)
            if not allow_ai_request(identity):
                raise HTTPException(status_code=429, detail="AI 分析请求过于频繁，请稍后重试")
        return generate_ai_analysis(hours=hours)
    except HTTPException:
        raise
    except Exception:
        return {
            "risk_level": "normal",
            "risk_score": 0,
            "summary": "AI 分析暂不可用",
            "key_findings": ["AI 分析服务暂不可用"],
            "trend_analysis": {"overall": "AI 分析暂不可用", "stations": []},
            "abnormal_stations": [],
            "recommendations": ["请稍后重试。"],
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "analysis_source": "rule_based",
            "note": "AI 分析暂不可用",
        }


if os.path.isdir(_dist_dir):
    _assets_dir = os.path.join(_dist_dir, "assets")

    @app.get("/{full_path:path}")
    def serve_spa(full_path: str):
        if full_path.startswith("api/"):
            raise HTTPException(status_code=404, detail="Not found")
        full_path = full_path.replace("\\", "/")
        if any(part == ".." for part in full_path.split("/")):
            raise HTTPException(status_code=404, detail="Not found")
        if full_path.startswith("assets/"):
            file_path = os.path.join(_dist_dir, full_path)
            if os.path.isfile(file_path):
                return FileResponse(file_path)
            raise HTTPException(status_code=404, detail="Not found")
        if full_path and os.path.isfile(os.path.join(_dist_dir, full_path)):
            return FileResponse(os.path.join(_dist_dir, full_path))
        return FileResponse(os.path.join(_dist_dir, "index.html"))
