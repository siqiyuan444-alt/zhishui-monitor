"""第16阶段：数据报表导出测试。"""

import io

import database
from fastapi.testclient import TestClient

import main

client = TestClient(main.app)


# ── /api/report ──────────────────────────────────────────────

def test_report_api_200():
    resp = client.get("/api/report?hours=24")
    assert resp.status_code == 200
    d = resp.json()
    assert d["hours"] == 24
    assert isinstance(d["count"], int)
    assert isinstance(d["data"], list)
    assert d["count"] >= 1


def test_report_api_with_station():
    resp = client.get("/api/report?station_id=ST001&hours=24")
    assert resp.status_code == 200
    d = resp.json()
    assert d["station_id"] == "ST001"
    assert d["station_name"] is not None
    for r in d["data"]:
        assert r["station_id"] == "ST001"
        assert "timestamp" in r
        assert "water_level" in r
        assert "warning_level" in r
        assert "rainfall" in r
        assert "status" in r
        assert "source" in r
        assert "data_quality" in r


def test_report_api_all_stations():
    resp = client.get("/api/report?hours=24")
    assert resp.status_code == 200
    d = resp.json()
    station_ids = {r["station_id"] for r in d["data"]}
    assert len(station_ids) >= 2
    assert d["station_id"] is None


def test_report_api_invalid_station_404():
    resp = client.get("/api/report?station_id=NOPE&hours=24")
    assert resp.status_code == 404


def test_report_api_hours_zero_422():
    resp = client.get("/api/report?hours=0")
    assert resp.status_code == 422


def test_report_api_hours_too_large_422():
    resp = client.get("/api/report?hours=99999")
    assert resp.status_code == 422


# ── /api/report/statistics ──────────────────────────────────

def test_report_statistics_200():
    resp = client.get("/api/report/statistics?hours=24")
    assert resp.status_code == 200
    d = resp.json()
    assert d["hours"] == 24
    assert isinstance(d["data_points"], int)
    assert isinstance(d["warning_count"], int)


def test_report_statistics_with_station():
    resp = client.get("/api/report/statistics?station_id=ST002&hours=24")
    assert resp.status_code == 200
    d = resp.json()
    assert d["station_id"] == "ST002"
    assert d["station_name"] is not None
    if d["data_points"] > 0:
        assert isinstance(d["avg_water_level"], (int, float))
        assert isinstance(d["max_water_level"], (int, float))
        assert isinstance(d["min_water_level"], (int, float))
        assert isinstance(d["avg_rainfall"], (int, float))
        assert isinstance(d["max_rainfall"], (int, float))
        assert d["max_water_level"] >= d["min_water_level"]


# ── CSV export ──────────────────────────────────────────────

def test_report_csv_200():
    resp = client.get("/api/report/export/csv?hours=24")
    assert resp.status_code == 200


def test_report_csv_content_type():
    resp = client.get("/api/report/export/csv?hours=24")
    assert "text/csv" in resp.headers["content-type"]


def test_report_csv_chinese_headers():
    resp = client.get("/api/report/export/csv?hours=24")
    content = resp.content.decode("utf-8")
    assert "站点编号" in content
    assert "站点名称" in content
    assert "采集时间" in content
    assert "当前水位" in content
    assert "警戒水位" in content
    assert "降雨量" in content
    assert "状态" in content
    assert "数据来源" in content
    assert "数据质量" in content


def test_report_csv_bom():
    resp = client.get("/api/report/export/csv?hours=24")
    assert resp.content[:3] == b"\xef\xbb\xbf"


def test_report_csv_empty_data():
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM water_data")
    conn.commit()
    conn.close()
    try:
        resp = client.get("/api/report/export/csv?hours=24")
        assert resp.status_code == 200
        content = resp.content.decode("utf-8")
        lines = content.strip().split("\n")
        assert len(lines) == 1
        assert "站点编号" in lines[0]
    finally:
        database._insert_initial_data(database.get_connection())
        conn = database.get_connection()
        database.ensure_history_backfill(conn)
        conn.close()


# ── Excel export ────────────────────────────────────────────

def test_report_excel_200():
    resp = client.get("/api/report/export/excel?hours=24")
    assert resp.status_code == 200


def test_report_excel_content_type():
    resp = client.get("/api/report/export/excel?hours=24")
    assert "spreadsheetml" in resp.headers["content-type"]


def test_report_excel_openable():
    from openpyxl import load_workbook

    resp = client.get("/api/report/export/excel?hours=24")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    assert ws["A1"].value is not None


def test_report_excel_chinese_headers():
    from openpyxl import load_workbook

    resp = client.get("/api/report/export/excel?hours=24")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    headers = [ws.cell(row=3, column=c).value for c in range(1, 10)]
    assert "站点编号" in headers
    assert "站点名称" in headers
    assert "采集时间" in headers
    assert "当前水位" in headers


def test_report_excel_empty_data():
    from openpyxl import load_workbook

    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM water_data")
    conn.commit()
    conn.close()
    try:
        resp = client.get("/api/report/export/excel?hours=24")
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws["A1"].value == "智慧水利 · 水情数据报表"
    finally:
        database._insert_initial_data(database.get_connection())
        conn = database.get_connection()
        database.ensure_history_backfill(conn)
        conn.close()
